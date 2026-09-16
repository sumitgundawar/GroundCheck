"""Reading documents into sections, and sections into chunks.

Supported formats: PDF, Word (.docx), HTML, Markdown and plain text. Each is
parsed into a title and a list of sections (a heading and its paragraphs), so
that every chunk keeps the heading it came from. Answers can then cite "Sepsis
protocol: Antibiotic timing" rather than an anonymous slice of text.

Headings come from the format's own structure where it has one (Word heading
styles, HTML h1 to h6, Markdown #). PDFs have no reliable structure, so a line
is treated as a heading when it is short, doesn't end like a sentence, and
looks like a title or a numbered heading.

Nothing here touches the database or the index; see app/knowledge.py."""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import PurePath


class DocumentError(ValueError):
    """The file can't be read. The message is safe to show."""


@dataclass
class Section:
    heading: str
    paragraphs: list[str] = field(default_factory=list)


@dataclass
class ParsedDocument:
    title: str
    sections: list[Section]
    media_type: str

    @property
    def text_length(self) -> int:
        return sum(len(p) for s in self.sections for p in s.paragraphs)


@dataclass(frozen=True)
class Chunk:
    position: int
    section: str
    text: str


MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".html": "text/html",
    ".htm": "text/html",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".txt": "text/plain",
}

MAX_FILE_BYTES = 25 * 1024 * 1024
_WS = re.compile(r"\s+")


def _clean(text: str) -> str:
    # NFKC turns typographic ligatures from PDFs (the single "fi" glyph) back
    # into plain letters, so keyword search and the coverage check can match.
    text = unicodedata.normalize("NFKC", text).replace("­", "")
    return _WS.sub(" ", text).strip()


def _add_paragraph(sections: list[Section], text: str) -> None:
    text = _clean(text)
    if not text:
        return
    if not sections:
        sections.append(Section(heading=""))
    sections[-1].paragraphs.append(text)


def _drop_empty(sections: list[Section]) -> list[Section]:
    return [s for s in sections if s.paragraphs]


# --- HTML ------------------------------------------------------------------

class _HTMLSections(HTMLParser):
    _BLOCK = {"p", "li", "td", "th", "dd", "dt", "blockquote", "pre", "caption", "figcaption", "div", "br", "tr"}
    _SKIP = {"script", "style", "nav", "footer", "header", "noscript", "template", "svg", "form"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title = ""
        self.sections: list[Section] = []
        self._buffer: list[str] = []
        self._heading: list[str] | None = None
        self._in_title = False
        self._skip_depth = 0

    def _flush(self) -> None:
        _add_paragraph(self.sections, "".join(self._buffer))
        self._buffer = []

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if tag == "title":
            self._in_title = True
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self._heading = []
        elif tag in self._BLOCK:
            self._flush()

    def handle_endtag(self, tag):
        if tag in self._SKIP:
            self._skip_depth = max(0, self._skip_depth - 1)
            return
        if self._skip_depth:
            return
        if tag == "title":
            self._in_title = False
        elif re.fullmatch(r"h[1-6]", tag) and self._heading is not None:
            heading = _clean("".join(self._heading))
            self._heading = None
            if heading:
                if tag == "h1" and not self.title:
                    self.title = heading
                self.sections.append(Section(heading=heading))
        elif tag in self._BLOCK:
            self._flush()

    def handle_data(self, data):
        if self._skip_depth:
            return
        if self._in_title:
            self.title = self.title or _clean(data)
        elif self._heading is not None:
            self._heading.append(data)
        else:
            self._buffer.append(data)

    def close(self):
        super().close()
        self._flush()


def parse_html(data: bytes) -> tuple[str, list[Section]]:
    parser = _HTMLSections()
    parser.feed(data.decode("utf-8", errors="replace"))
    parser.close()
    return parser.title, _drop_empty(parser.sections)


# --- Word --------------------------------------------------------------------

def parse_docx(data: bytes) -> tuple[str, list[Section]]:
    try:
        import docx
        from docx.document import Document as DocxDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - dependency is in requirements
        raise DocumentError("Word support needs python-docx.") from exc
    try:
        document: DocxDocument = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - any failure means an unreadable file
        raise DocumentError("This Word file couldn't be opened. Save it as .docx and try again.") from exc

    title = _clean(document.core_properties.title or "")
    sections: list[Section] = []
    # Walk paragraphs and tables in document order.
    for block in document.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(block, document)
            style = (paragraph.style.name if paragraph.style is not None else "").lower()
            text = _clean(paragraph.text)
            if not text:
                continue
            if style == "title":
                title = title or text
            elif style.startswith("heading"):
                sections.append(Section(heading=text))
            else:
                _add_paragraph(sections, text)
        elif tag == "tbl":
            for row in Table(block, document).rows:
                cells = [_clean(c.text) for c in row.cells]
                _add_paragraph(sections, " | ".join(dict.fromkeys(c for c in cells if c)))
    return title, _drop_empty(sections)


# --- PDF ---------------------------------------------------------------------

_NUMBERED_HEADING = re.compile(r"^(\d+(\.\d+)*\.?|[A-Z]\.)\s+\S")


def _looks_like_heading(line: str) -> bool:
    if not 2 <= len(line) <= 80 or line.endswith((".", ",", ";", ":")):
        return False
    words = line.split()
    if len(words) > 10 or not re.search(r"[A-Za-z]", line):
        return False
    if _NUMBERED_HEADING.match(line):
        return True
    if line.isupper() and len(words) <= 8:
        return True
    capitalised = sum(1 for w in words if w[0].isupper() or not w[0].isalpha())
    return len(words) <= 6 and capitalised == len(words)


def parse_pdf(data: bytes) -> tuple[str, list[Section]]:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("PDF support needs pypdf.") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError("This PDF is password-protected. Remove the password and try again.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except DocumentError:
        raise
    except (PdfReadError, Exception) as exc:  # noqa: BLE001
        raise DocumentError("This PDF couldn't be read.") from exc

    title = _clean(str((reader.metadata or {}).get("/Title", "") or ""))
    sections: list[Section] = []
    paragraph: list[str] = []

    def flush():
        if paragraph:
            _add_paragraph(sections, " ".join(paragraph))
            paragraph.clear()

    for page in pages:
        for raw in page.splitlines():
            line = _clean(raw)
            if not line:
                flush()
                continue
            if _looks_like_heading(line) and (not paragraph or paragraph[-1].endswith((".", ":", "!", "?"))):
                flush()
                if not title and not sections:
                    title = line
                    continue
                sections.append(Section(heading=line))
                continue
            # Join hyphenated line breaks: "anti-\nbiotic" -> "antibiotic".
            if paragraph and paragraph[-1].endswith("-") and line[:1].islower():
                paragraph[-1] = paragraph[-1][:-1] + line
            else:
                paragraph.append(line)
            if line.endswith((".", "!", "?")) and len(" ".join(paragraph)) > 400:
                flush()
        flush()

    if not any(s.paragraphs for s in sections) and sum(len(p) for p in pages) < 20:
        raise DocumentError("This PDF has no extractable text. It may be a scan; run it through OCR first.")
    return title, _drop_empty(sections)


# --- Markdown and text -----------------------------------------------------

def parse_text(data: bytes, markdown: bool) -> tuple[str, list[Section]]:
    text = data.decode("utf-8", errors="replace")
    title = ""
    sections: list[Section] = []
    paragraph: list[str] = []

    def flush():
        if paragraph:
            _add_paragraph(sections, " ".join(paragraph))
            paragraph.clear()

    for raw in text.splitlines():
        line = raw.strip()
        heading = re.match(r"^(#{1,6})\s+(.*)$", line) if markdown else None
        if heading:
            flush()
            name = _clean(heading.group(2).strip("#"))
            if len(heading.group(1)) == 1 and not title:
                title = name
            sections.append(Section(heading=name))
        elif not line:
            flush()
        else:
            paragraph.append(re.sub(r"^[-*+]\s+|^\d+[.)]\s+", "", line) if markdown else line)
    flush()
    return title, _drop_empty(sections)


# --- Entry point -----------------------------------------------------------

def parse(filename: str, data: bytes) -> ParsedDocument:
    suffix = PurePath(filename or "").suffix.lower()
    if suffix not in MEDIA_TYPES:
        raise DocumentError("Upload a PDF, Word (.docx), HTML, Markdown or text file.")
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError(f"Files can be up to {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    if not data:
        raise DocumentError("The file is empty.")

    if suffix == ".pdf":
        title, sections = parse_pdf(data)
    elif suffix == ".docx":
        title, sections = parse_docx(data)
    elif suffix in (".html", ".htm"):
        title, sections = parse_html(data)
    else:
        title, sections = parse_text(data, markdown=suffix in (".md", ".markdown"))

    if not sections:
        raise DocumentError("No readable text was found in this file.")
    fallback = PurePath(filename).stem.replace("_", " ").replace("-", " ").strip()
    return ParsedDocument(title=title or fallback or "Untitled document", sections=sections,
                          media_type=MEDIA_TYPES[suffix])


_SENTENCE_END = re.compile(r"(?<=[.!?])\s+")


def chunk(document: ParsedDocument, max_chars: int = 1000) -> list[Chunk]:
    """Split each section into chunks of whole paragraphs, up to max_chars.
    A paragraph longer than that is split at sentence boundaries. Chunks never
    cross a section boundary, so each keeps its heading."""
    chunks: list[Chunk] = []

    def emit(section: str, text: str) -> None:
        text = text.strip()
        if text:
            chunks.append(Chunk(position=len(chunks), section=section, text=text))

    for section in document.sections:
        heading = section.heading or document.title
        current = ""
        pieces: list[str] = []
        for paragraph in section.paragraphs:
            if len(paragraph) <= max_chars:
                pieces.append(paragraph)
                continue
            sentence_buffer = ""
            for sentence in _SENTENCE_END.split(paragraph):
                while len(sentence) > max_chars:  # a single enormous "sentence"
                    pieces.append(sentence[:max_chars])
                    sentence = sentence[max_chars:]
                if sentence_buffer and len(sentence_buffer) + 1 + len(sentence) > max_chars:
                    pieces.append(sentence_buffer)
                    sentence_buffer = sentence
                else:
                    sentence_buffer = f"{sentence_buffer} {sentence}".strip()
            if sentence_buffer:
                pieces.append(sentence_buffer)

        for piece in pieces:
            if current and len(current) + 1 + len(piece) > max_chars:
                emit(heading, current)
                current = piece
            else:
                current = f"{current} {piece}".strip()
        emit(heading, current)
    return chunks
