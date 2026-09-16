"""Organisation documents: parsing, chunking, versions, approval, effective
dates, index rebuilds, per-document evaluation, and the API's rules.

Index tests build a small, isolated index in a temp folder (demo corpus
excluded) and restore the real one afterwards."""

from __future__ import annotations

import io
import os
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest

os.environ["GROQ_API_KEY"] = ""
os.environ["FORCE_EXTRACTIVE"] = "true"
os.environ["RATE_LIMIT_PER_MINUTE"] = "1000000"
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient  # noqa: E402

from app import auth, config, db, documents, knowledge, pipeline, retrieval  # noqa: E402

FIXTURES = Path(__file__).parent / "fixtures"
PDF = FIXTURES / "neutropenic_fever_protocol.pdf"
PASSWORD = "correct horse battery staple"

HTML = b"""<!doctype html><html><head><title>Falls Prevention Policy</title>
<script>var tracking = 1;</script></head><body><nav>Home | Policies</nav>
<h1>Falls Prevention Policy</h1><p>This fictional policy is for import testing.</p>
<h2>Risk assessment</h2><p>Complete a falls risk assessment within 6 hours of admission.</p>
<ul><li>Check footwear.</li><li>Check eyesight &amp; hearing aids.</li></ul>
<h2>After a fall</h2><p>Record the fall and review medication within 24 hours.</p>
<footer>Footer text</footer></body></html>"""


def _docx_bytes() -> bytes:
    import docx

    d = docx.Document()
    d.core_properties.title = "Blood Transfusion Checklist"
    d.add_paragraph("Blood Transfusion Checklist", style="Title")
    d.add_heading("Before starting", level=1)
    d.add_paragraph("Confirm patient identity at the bedside with two staff members.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Observation", "When"
    table.cell(1, 0).text, table.cell(1, 1).text = "Temperature", "Before and at 15 minutes"
    d.add_heading("During the transfusion", level=1)
    d.add_paragraph("Stop the transfusion if the patient develops a fever or rash.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# --- Parsing -----------------------------------------------------------------

def test_pdf_sections_and_ligatures():
    doc = documents.parse(PDF.name, PDF.read_bytes())
    assert doc.title == "Neutropenic Fever Protocol"
    headings = [s.heading for s in doc.sections]
    assert headings[1:] == ["1. Recognition", "2. First Hour Actions", "3. Review"]
    body = " ".join(p for s in doc.sections for p in s.paragraphs)
    assert "first dose of Zorbaxil 2 g" in body
    assert "ﬁ" not in body  # the "fi" ligature is normalised


def test_html_skips_navigation_scripts_and_footers():
    doc = documents.parse("falls.html", HTML)
    assert doc.title == "Falls Prevention Policy"
    assert [s.heading for s in doc.sections] == ["Falls Prevention Policy", "Risk assessment", "After a fall"]
    body = " ".join(p for s in doc.sections for p in s.paragraphs)
    assert "Check eyesight & hearing aids." in body
    assert "tracking" not in body and "Home | Policies" not in body and "Footer text" not in body


def test_docx_headings_paragraphs_and_tables():
    doc = documents.parse("checklist.docx", _docx_bytes())
    assert doc.title == "Blood Transfusion Checklist"
    assert [s.heading for s in doc.sections] == ["Before starting", "During the transfusion"]
    assert "Temperature | Before and at 15 minutes" in doc.sections[0].paragraphs


def test_markdown_and_text():
    md = documents.parse("guide.md", b"# Oxygen Guide\n\nIntro text.\n\n## Targets\n\n- Aim for 94 to 98 percent.\n")
    assert md.title == "Oxygen Guide"
    assert md.sections[-1].heading == "Targets"
    assert md.sections[-1].paragraphs == ["Aim for 94 to 98 percent."]
    txt = documents.parse("notes_on_pain.txt", b"Line one.\n\nLine two.")
    assert txt.title == "notes on pain"


@pytest.mark.parametrize("name, data, message", [
    ("virus.exe", b"MZ", "Upload a PDF"),
    ("empty.txt", b"", "empty"),
    ("broken.pdf", b"%PDF-1.4 not really", "couldn't be read"),
    ("broken.docx", b"not a zip", "couldn't be opened"),
    ("blank.html", b"<html><body><script>x</script></body></html>", "No readable text"),
])
def test_unreadable_files_give_clear_errors(name, data, message):
    with pytest.raises(documents.DocumentError, match=message):
        documents.parse(name, data)


def test_chunks_respect_size_and_sections():
    long_paragraph = " ".join(f"Sentence number {i} is here." for i in range(200))
    doc = documents.ParsedDocument("Long", [documents.Section("A", [long_paragraph]),
                                            documents.Section("B", ["Short."])], "text/plain")
    chunks = documents.chunk(doc, max_chars=300)
    assert all(len(c.text) <= 300 for c in chunks)
    assert chunks[-1].section == "B" and chunks[-1].text == "Short."
    assert [c.position for c in chunks] == list(range(len(chunks)))
    rebuilt = " ".join(c.text for c in chunks if c.section == "A")
    assert rebuilt == long_paragraph


# --- Lifecycle -------------------------------------------------------------



def test_import_is_pending_with_chunks(database):
    source = knowledge.import_document(PDF.name, PDF.read_bytes(), owner="Oncology")
    assert source["status"] == "pending" and source["version"] == 1
    assert source["chunks"] == 4 and not source["in_force"]
    detail = knowledge.get_source(source["id"])
    assert detail["sections"][2]["chunk_id"] == f"DOC{source['id']}-003"
    assert knowledge.approved_records() == []


def test_duplicate_file_is_refused_and_new_content_is_a_new_version(database):
    first = knowledge.import_document("falls.html", HTML)
    with pytest.raises(knowledge.KnowledgeError, match="already version 1"):
        knowledge.import_document("falls.html", HTML)
    second = knowledge.import_document("falls.html", HTML.replace(b"6 hours", b"4 hours"))
    assert second["document_key"] == first["document_key"] and second["version"] == 2


def test_uploader_cannot_approve_their_own_document(database):
    uploader = auth.create_user("uploader@example.org", PASSWORD, role="admin")
    reviewer = auth.create_user("reviewer@example.org", PASSWORD, role="admin")
    source = knowledge.import_document("falls.html", HTML, uploaded_by=uploader.id)
    with pytest.raises(knowledge.KnowledgeError, match="other than the uploader"):
        knowledge.review(source["id"], "approve", uploader.id)
    approved = knowledge.review(source["id"], "approve", reviewer.id, note="Checked against the ward copy.")
    assert approved["status"] == "approved" and approved["in_force"]


def test_approving_a_new_version_retires_the_old_one(database):
    v1 = knowledge.import_document("falls.html", HTML)
    knowledge.review(v1["id"], "approve", None, allow_self_approval=True)
    v2 = knowledge.import_document("falls.html", HTML.replace(b"6 hours", b"4 hours"))
    knowledge.review(v2["id"], "approve", None, allow_self_approval=True)
    statuses = {s["version"]: s["status"] for s in knowledge.list_sources()}
    assert statuses == {1: "retired", 2: "approved"}
    texts = " ".join(r["text"] for r in knowledge.approved_records())
    assert "4 hours" in texts and "6 hours" not in texts


def test_review_rules(database):
    source = knowledge.import_document("falls.html", HTML)
    with pytest.raises(knowledge.KnowledgeError, match="Only approved"):
        knowledge.review(source["id"], "retire", None)
    knowledge.review(source["id"], "reject", None, note="Out of date")
    with pytest.raises(knowledge.KnowledgeError, match="Only pending"):
        knowledge.review(source["id"], "approve", None, allow_self_approval=True)
    assert knowledge.get_source(source["id"])["review_note"] == "Out of date"


def test_effective_and_expiry_dates(database):
    now = datetime.now(timezone.utc)
    with pytest.raises(knowledge.KnowledgeError, match="after the effective date"):
        knowledge.import_document("falls.html", HTML, effective_from=now, expires_on=now - timedelta(days=1))
    future = knowledge.import_document("falls.html", HTML, effective_from=now + timedelta(days=7))
    knowledge.review(future["id"], "approve", None, allow_self_approval=True)
    assert not knowledge.get_source(future["id"])["in_force"]


# --- Index and answers ---------------------------------------------------------

@pytest.fixture()
def isolated_index(database, monkeypatch, tmp_path):
    monkeypatch.setattr(config, "INDEX_DIR", tmp_path / "index")
    monkeypatch.setattr(config, "INCLUDE_DEMO_CORPUS", False)
    monkeypatch.setattr(config, "VECTOR_STORE", "local")
    yield
    monkeypatch.undo()
    db.reset_engine()
    retrieval.load_index()  # restore the real index for other tests


def _approve(filename: str, data: bytes, **kwargs) -> dict:
    source = knowledge.import_document(filename, data, **kwargs)
    return knowledge.review(source["id"], "approve", None, allow_self_approval=True)


def test_approved_documents_are_answered_and_cited(isolated_index):
    source = _approve(PDF.name, PDF.read_bytes())
    status = knowledge.rebuild_index()
    assert status["documents"] == 4 and status["embedded"] == 4
    assert knowledge.rebuild_index()["reused"] == 4  # cached embeddings

    r = pipeline.run("When should the first dose of Zorbaxil be given for neutropenic fever?")
    assert r.decision == "answer"
    assert any(sid.startswith(f"DOC{source['id']}-") for c in r.claims for sid in c.source_ids)

    refused = pipeline.run("What is the first dose of Quorbaxil for neutropenic fever?")
    assert refused.decision == "refuse"


def test_retired_and_expired_documents_are_not_cited(isolated_index):
    now = datetime.now(timezone.utc)
    question = "How soon should a falls risk assessment be completed after admission?"
    falls = _approve("falls.html", HTML, expires_on=now + timedelta(days=30))
    knowledge.rebuild_index()
    assert pipeline.run(question).decision == "answer"

    # An expired document stays in the index but is never retrieved.
    with db.session() as s:
        s.get(db.Source, falls["id"]).expires_on = now - timedelta(minutes=1)
    knowledge.rebuild_index()
    assert len(retrieval.all_metadata()) > 0
    assert pipeline.run(question).decision == "refuse"

    # A retired document leaves the index entirely.
    fresh = _approve("falls.html", HTML.replace(b"6 hours", b"4 hours"))
    knowledge.rebuild_index()
    assert pipeline.run(question).decision == "answer"
    knowledge.review(fresh["id"], "retire", None)
    knowledge.rebuild_index()
    assert all(r.get("source_id") != fresh["id"] for r in retrieval.all_metadata())
    assert pipeline.run(question).decision == "refuse"


def test_an_empty_index_refuses_everything(isolated_index):
    status = knowledge.rebuild_index()
    assert status["documents"] == 0
    r = pipeline.run("What is the standard dose of Caloradine?")
    assert r.decision == "refuse"
    assert retrieval.corpus_projection.__wrapped__() == []


def test_search_skips_rows_outside_their_dates(isolated_index):
    now = datetime.now(timezone.utc)
    _approve("falls.html", HTML, effective_from=now - timedelta(days=1), expires_on=now + timedelta(days=1))
    knowledge.rebuild_index()
    question = "How soon should a falls risk assessment be completed after admission?"
    assert retrieval.search(question, 4)
    retrieval._not_after[:] = (now - timedelta(seconds=1)).timestamp()
    assert retrieval.search(question, 4) == []


def test_document_evaluation(isolated_index):
    source = _approve(PDF.name, PDF.read_bytes())
    knowledge.rebuild_index()
    result = knowledge.evaluate_source(source["id"])
    assert result["total"] == 8
    assert result["unsafe_answers"] == 0
    assert result["passed"] >= 6
    assert knowledge.get_source(source["id"])["evaluation"]["total"] == 8


# --- API -----------------------------------------------------------------------

@pytest.fixture()
def client(isolated_index, monkeypatch):
    knowledge.rebuild_index()  # the app loads the index when it starts
    monkeypatch.setattr(config, "AUTH_REQUIRED", True)
    monkeypatch.setattr(config, "SESSION_COOKIE_SECURE", False)
    from app.main import app
    with TestClient(app) as c:
        yield c


def _login(client, email):
    client.post("/api/auth/logout")
    assert client.post("/api/auth/login", json={"email": email, "password": PASSWORD}).status_code == 200


def test_api_upload_review_and_rebuild(client):
    auth.create_user("clin@example.org", PASSWORD, role="clinician")
    auth.create_user("rev@example.org", PASSWORD, role="reviewer")
    auth.create_user("admin1@example.org", PASSWORD, role="admin")
    auth.create_user("admin2@example.org", PASSWORD, role="admin")

    _login(client, "clin@example.org")
    files = {"file": (PDF.name, PDF.read_bytes(), "application/pdf")}
    assert client.post("/api/sources", files=files).status_code == 403

    _login(client, "admin1@example.org")
    r = client.post("/api/sources", files=files, data={"owner": "Oncology", "expires_on": "2099-01-01"})
    assert r.status_code == 200, r.text
    source_id = r.json()["source"]["id"]
    bad = client.post("/api/sources", files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert bad.status_code == 400 and "Upload a PDF" in bad.json()["detail"]
    self_approval = client.post(f"/api/sources/{source_id}/approve", json={"note": ""})
    assert self_approval.status_code == 400 and "uploader" in self_approval.json()["detail"]

    _login(client, "rev@example.org")
    assert client.get("/api/sources").status_code == 200
    assert client.post(f"/api/sources/{source_id}/approve", json={}).status_code == 403

    _login(client, "admin2@example.org")
    approved = client.post(f"/api/sources/{source_id}/approve", json={"note": "Matches the signed copy."})
    assert approved.status_code == 200 and approved.json()["source"]["status"] == "approved"
    knowledge._index_lock.acquire()  # wait for the background rebuild to finish
    knowledge._index_lock.release()
    detail = client.get(f"/api/sources/{source_id}").json()["source"]
    assert detail["reviewed_by"] is not None and len(detail["sections"]) == 4
    assert client.post("/api/sources/999/nonsense", json={}).status_code == 404
